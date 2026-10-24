from __future__ import annotations

from io import BytesIO

from docx import Document

from office_api.docx import fill_document


def template_bytes(build) -> bytes:
    document = Document()
    build(document)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def document_text(content: bytes) -> str:
    document = Document(BytesIO(content))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    for section in document.sections:
        for story in (section.header, section.footer):
            parts.extend(paragraph.text for paragraph in story.paragraphs)
    return "\n".join(part for part in parts if part)


def test_fills_body_table_header_and_footer() -> None:
    def build(document) -> None:
        document.add_paragraph("Site: {{site}}")
        table = document.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "Result: {{value}}"
        document.sections[0].header.paragraphs[0].text = "{{title}}"
        document.sections[0].footer.paragraphs[0].text = "{{site}}"

    filled, unfilled = fill_document(
        template_bytes(build), {"site": "Plant A", "value": 42, "title": "Report"}
    )

    text = document_text(filled)
    assert "Site: Plant A" in text
    assert "Result: 42" in text
    assert "Report" in text
    assert unfilled == []


def test_placeholder_split_across_runs_is_still_filled() -> None:
    def build(document) -> None:
        paragraph = document.add_paragraph()
        paragraph.add_run("Site: {{si")
        paragraph.add_run("te}}")

    filled, unfilled = fill_document(template_bytes(build), {"site": "Plant B"})

    assert document_text(filled) == "Site: Plant B"
    assert unfilled == []


def test_unknown_placeholder_is_empty_and_reported() -> None:
    def build(document) -> None:
        document.add_paragraph("A={{a}} B={{missing}}")

    filled, unfilled = fill_document(template_bytes(build), {"a": "1"})

    assert document_text(filled) == "A=1 B="
    assert unfilled == ["missing"]


def test_input_bytes_are_not_modified() -> None:
    def build(document) -> None:
        document.add_paragraph("{{name}}")

    template = template_bytes(build)
    snapshot = bytes(template)
    fill_document(template, {"name": "x"})

    assert template == snapshot

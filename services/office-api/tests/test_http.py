from __future__ import annotations

import hashlib
import json
from io import BytesIO

import pytest
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook

from office_api import main


@pytest.fixture()
def client() -> TestClient:
    return TestClient(main.app)


def xlsx_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Monitoring"
    sheet.append(["pollutant", "conc_mg_l"])
    sheet.append(["COD", 42.5])
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def docx_bytes() -> bytes:
    document = Document()
    document.add_paragraph("Site: {{site}} {{missing}}")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_health_and_ready(client: TestClient) -> None:
    assert client.get("/healthz").json()["status"] == "ok"
    assert client.get("/readyz").json()["status"] == "ready"


def test_parse_endpoint_returns_tables(client: TestClient) -> None:
    response = client.post(
        "/v1/xlsx/parse",
        files={"file": ("monitoring.xlsx", xlsx_bytes())},
    )

    assert response.status_code == 200
    sheets = response.json()["sheets"]
    assert sheets[0]["name"] == "Monitoring"
    assert sheets[0]["tables"][0]["rows"] == [["COD", 42.5]]


def test_fill_endpoint_returns_docx_with_headers(client: TestClient) -> None:
    response = client.post(
        "/v1/docx/fill",
        files={"file": ("template.docx", docx_bytes())},
        data={"values": json.dumps({"site": "Plant A"})},
    )

    assert response.status_code == 200
    body = response.content
    assert response.headers["content-type"].startswith(main.DOCX_MIME)
    assert response.headers["X-Office-Doc-SHA256"] == hashlib.sha256(body).hexdigest()
    assert response.headers["X-Office-Unfilled"] == "missing"
    filled = Document(BytesIO(body))
    assert filled.paragraphs[0].text == "Site: Plant A "


def test_malformed_xlsx_is_422(client: TestClient) -> None:
    response = client.post(
        "/v1/xlsx/parse",
        files={"file": ("broken.xlsx", b"not a zip")},
    )

    assert response.status_code == 422
    assert "detail" in response.json()


def test_malformed_docx_is_422(client: TestClient) -> None:
    response = client.post(
        "/v1/docx/fill",
        files={"file": ("broken.docx", b"not a zip")},
        data={"values": "{}"},
    )

    assert response.status_code == 422


def test_values_must_be_a_json_object(client: TestClient) -> None:
    response = client.post(
        "/v1/docx/fill",
        files={"file": ("template.docx", docx_bytes())},
        data={"values": "[1, 2]"},
    )

    assert response.status_code == 422


def test_oversized_upload_is_413(client: TestClient, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(main, "MAX_BYTES", 8)
    response = client.post(
        "/v1/xlsx/parse",
        files={"file": ("big.xlsx", xlsx_bytes())},
    )

    assert response.status_code == 413

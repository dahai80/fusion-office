# Parses .xlsx workbooks into the tables described by issue #1. Sheets are split
# into blocks on empty rows and columns; blocks without a usable header land in
# unparsed_regions rather than being guessed.

from __future__ import annotations

import datetime as dt
import logging
from io import BytesIO
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

log = logging.getLogger(__name__)

CellValue = str | int | float | bool | None


def _normalize(value: Any) -> CellValue:
    if value is None:
        return None
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return value.isoformat()
    if isinstance(value, (int, float, bool, str)):
        return value
    return str(value)


def _text(value: CellValue) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _grid(sheet: Worksheet) -> tuple[list[list[CellValue]], int, int]:
    rows = max(sheet.max_row or 0, 1)
    cols = max(sheet.max_column or 0, 1)
    for merged in sheet.merged_cells.ranges:
        rows = max(rows, merged.max_row)
        cols = max(cols, merged.max_col)
    grid: list[list[CellValue]] = [[None for _ in range(cols)] for _ in range(rows)]
    for row in range(rows):
        for col in range(cols):
            grid[row][col] = _normalize(sheet.cell(row=row + 1, column=col + 1).value)
    for merged in sheet.merged_cells.ranges:
        anchor = grid[merged.min_row - 1][merged.min_col - 1]
        for row in range(merged.min_row, merged.max_row + 1):
            for col in range(merged.min_col, merged.max_col + 1):
                grid[row - 1][col - 1] = anchor
    return grid, rows, cols


def _occupied(grid: list[list[CellValue]], row: int, col: int) -> bool:
    return _text(grid[row][col]) != ""


def _bands(grid: list[list[CellValue]], rows: int, cols: int) -> list[tuple[int, int]]:
    bands: list[tuple[int, int]] = []
    start: int | None = None
    for row in range(rows):
        filled = any(_occupied(grid, row, col) for col in range(cols))
        if filled and start is None:
            start = row
        elif not filled and start is not None:
            bands.append((start, row - 1))
            start = None
    if start is not None:
        bands.append((start, rows - 1))
    return bands


def _column_runs(
    grid: list[list[CellValue]], top: int, bottom: int, cols: int
) -> list[tuple[int, int]]:
    runs: list[tuple[int, int]] = []
    start: int | None = None
    for col in range(cols):
        filled = any(_occupied(grid, row, col) for row in range(top, bottom + 1))
        if filled and start is None:
            start = col
        elif not filled and start is not None:
            runs.append((start, col - 1))
            start = None
    if start is not None:
        runs.append((start, cols - 1))
    return runs


def _range_name(top: int, left: int, bottom: int, right: int) -> str:
    return f"{get_column_letter(left + 1)}{top + 1}:{get_column_letter(right + 1)}{bottom + 1}"


def _merged_in_block(
    sheet: Worksheet, top: int, left: int, bottom: int, right: int
) -> list[list[str]]:
    found: list[list[str]] = []
    for merged in sorted(sheet.merged_cells.ranges, key=lambda r: (r.min_row, r.min_col)):
        if (
            merged.min_row - 1 >= top
            and merged.max_row - 1 <= bottom
            and merged.min_col - 1 >= left
            and merged.max_col - 1 <= right
        ):
            found.append(
                [
                    f"r{merged.min_row}c{merged.min_col}",
                    f"r{merged.max_row}c{merged.max_col}",
                ]
            )
    return found


def _table_from_block(
    sheet: Worksheet,
    grid: list[list[CellValue]],
    top: int,
    left: int,
    bottom: int,
    right: int,
) -> dict[str, Any] | None:
    if bottom <= top:
        return None
    headers = [_text(grid[top][col]) for col in range(left, right + 1)]
    if not headers or any(not header for header in headers):
        return None
    if len(set(headers)) != len(headers):
        log.info(
            "sheet %s block %s has duplicate headers",
            sheet.title,
            _range_name(top, left, bottom, right),
        )
        return None
    rows: list[list[CellValue]] = []
    for row in range(top + 1, bottom + 1):
        rows.append([grid[row][col] for col in range(left, right + 1)])
    if not rows:
        return None
    return {
        "header_row": 0,
        "headers": headers,
        "rows": rows,
        "merged_cells": _merged_in_block(sheet, top, left, bottom, right),
        "range": _range_name(top, left, bottom, right),
    }


def parse_workbook(content: bytes) -> dict[str, Any]:
    workbook = load_workbook(BytesIO(content), data_only=True)
    sheets: list[dict[str, Any]] = []
    try:
        for sheet in workbook.worksheets:
            grid, rows, cols = _grid(sheet)
            tables: list[dict[str, Any]] = []
            unparsed: list[str] = []
            for top, bottom in _bands(grid, rows, cols):
                for left, right in _column_runs(grid, top, bottom, cols):
                    block = _range_name(top, left, bottom, right)
                    table = _table_from_block(sheet, grid, top, left, bottom, right)
                    if table is None:
                        unparsed.append(block)
                    else:
                        tables.append(table)
            sheets.append({"name": sheet.title, "tables": tables, "unparsed_regions": unparsed})
            log.info(
                "sheet %s: %d tables, %d unparsed regions",
                sheet.title,
                len(tables),
                len(unparsed),
            )
    finally:
        workbook.close()
    return {"sheets": sheets}

# Fills {{placeholder}} markers in a .docx template (issue #1, G3). Markers are
# replaced across run boundaries, and unknown placeholders are reported back.

from __future__ import annotations

import logging
import re
from collections.abc import Iterable
from io import BytesIO
from typing import Any

from docx import Document
from docx.document import Document as DocxDocument
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph

log = logging.getLogger(__name__)

PLACEHOLDER = re.compile(r"\{\{\s*([^{}]+?)\s*\}\}")


def _replace_span(runs: list[Any], start: int, end: int, replacement: str) -> None:
    offset = 0
    first = None
    first_offset = 0
    last = None
    last_offset = 0
    for index, run in enumerate(runs):
        length = len(run.text)
        if first is None and start < offset + length:
            first = index
            first_offset = start - offset
        if end <= offset + length:
            last = index
            last_offset = end - offset
            break
        offset += length
    if first is None or last is None:
        return
    if first == last:
        text = runs[first].text
        runs[first].text = text[:first_offset] + replacement + text[last_offset:]
        return
    head = runs[first].text[:first_offset]
    tail = runs[last].text[last_offset:]
    runs[first].text = head + replacement
    for index in range(first + 1, last):
        runs[index].text = ""
    runs[last].text = tail


def _fill_paragraph(paragraph: Paragraph, values: dict[str, Any], found: set[str]) -> None:
    runs = paragraph.runs
    if not runs:
        return
    text = "".join(run.text for run in runs)
    matches = list(PLACEHOLDER.finditer(text))
    if not matches:
        return
    for match in reversed(matches):
        name = match.group(1).strip()
        found.add(name)
        _replace_span(runs, match.start(), match.end(), str(values.get(name, "")))


def _fill_tables(tables: Iterable[Table], values: dict[str, Any], found: set[str]) -> None:
    for table in tables:
        for row in table.rows:
            for cell in row.cells:
                _fill_cell(cell, values, found)


def _fill_cell(cell: _Cell, values: dict[str, Any], found: set[str]) -> None:
    for paragraph in cell.paragraphs:
        _fill_paragraph(paragraph, values, found)
    _fill_tables(cell.tables, values, found)


def _fill_story(story: Any, values: dict[str, Any], found: set[str]) -> None:
    for paragraph in story.paragraphs:
        _fill_paragraph(paragraph, values, found)
    _fill_tables(story.tables, values, found)


def fill_document(content: bytes, values: dict[str, Any]) -> tuple[bytes, list[str]]:
    document: DocxDocument = Document(BytesIO(content))
    found: set[str] = set()
    _fill_story(document, values, found)
    for section in document.sections:
        for story in (
            section.header,
            section.footer,
            section.first_page_header,
            section.first_page_footer,
            section.even_page_header,
            section.even_page_footer,
        ):
            _fill_story(story, values, found)
    output = BytesIO()
    document.save(output)
    filled = output.getvalue()
    unfilled = sorted(name for name in found if name not in values)
    log.info(
        "filled docx: %d placeholders, %d unfilled, %d bytes out",
        len(found),
        len(unfilled),
        len(filled),
    )
    return filled, unfilled
